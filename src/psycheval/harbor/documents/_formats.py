"""Format adapters with deferred optional imports."""

import io

from ._limits import DocumentParseError, InspectionBudget
from .models import (
    DocxReport,
    Paragraph,
    PdfPage,
    PdfReport,
    Table,
    TableCell,
    TextRun,
    XlsxCell,
    XlsxReport,
    XlsxSheet,
)


def inspect_pdf(snapshot, backend, limits):
    from pypdf.errors import DependencyError, PyPdfError

    from . import MissingDocumentDependency

    budget = InspectionBudget(limits)
    try:
        reader = backend.PdfReader(io.BytesIO(snapshot.content), strict=True)
        if reader.is_encrypted:
            raise DocumentParseError("Encrypted PDF is unsupported")
        budget.pages(len(reader.pages))
        pages = []
        for index, page in enumerate(reader.pages, 1):
            budget.fact()
            text = page.extract_text() or ""
            budget.text(text)
            pages.append(
                PdfPage(
                    index,
                    text,
                    float(page.mediabox.width),
                    float(page.mediabox.height),
                )
            )
    except DependencyError as error:
        raise MissingDocumentDependency(f"PDF backend dependency: {error}") from error
    except PyPdfError as error:
        raise DocumentParseError(f"Cannot read PDF: {error}") from error
    return PdfReport(snapshot, tuple(pages), limits.text_bytes)


def inspect_xlsx(snapshot, backend, limits):
    from openpyxl.utils.exceptions import InvalidFileException

    # openpyxl exposes either formulas or their stored values per load.
    # Both reads use the same immutable bytes; no recalculation is performed.
    def load(data_only):
        try:
            return backend.load_workbook(
                io.BytesIO(snapshot.content), data_only=data_only, keep_links=False
            )
        except (OSError, InvalidFileException) as error:
            # These reads are from BytesIO. openpyxl uses OSError for missing
            # workbook parts; source filesystem errors occur before this boundary.
            raise DocumentParseError(f"Cannot read XLSX: {error}") from error

    workbook = load(False)
    cached = None
    budget = InspectionBudget(limits)
    try:
        budget.pages(len(workbook.worksheets))
        cached = load(True)
        sheets = []
        for sheet in workbook.worksheets:
            budget.text(sheet.title)
            cached_sheet = cached[sheet.title]
            cells = []
            # Iterate stored cells: a distant cell must not expand a sparse sheet
            # into billions of newly allocated empty cells via iter_rows().
            # Check before sorting/materializing the sparse index.
            budget.fact(len(sheet._cells))
            for _, cell in sorted(sheet._cells.items()):
                if cell.value is None:
                    continue
                formula = None
                if cell.data_type == "f":
                    formula = (
                        cell.value
                        if isinstance(cell.value, str)
                        else getattr(cell.value, "text", None)
                    )
                    if formula is None:
                        raise ValueError(
                            f"Unsupported formula at {sheet.title}!{cell.coordinate}"
                        )
                value = None if formula else cell.value
                cached_value = cached_sheet[cell.coordinate].value if formula else None
                for text in (value, formula, cached_value):
                    budget.text(text)
                cells.append(
                    XlsxCell(
                        cell.coordinate,
                        value,
                        formula,
                        cached_value,
                    )
                )
            budget.fact(len(sheet.merged_cells.ranges))
            merged_ranges = tuple(str(item) for item in sheet.merged_cells.ranges)
            for merged in merged_ranges:
                budget.text(merged)
            sheets.append(
                XlsxSheet(
                    sheet.title,
                    sheet.sheet_state,
                    tuple(cells),
                    merged_ranges,
                )
            )
        return XlsxReport(snapshot, tuple(sheets), limits.text_bytes)
    finally:
        workbook.close()
        if cached is not None:
            cached.close()


def inspect_docx(snapshot, backend, limits):
    from docx.opc.exceptions import OpcError
    from docx.table import Table as BackendTable
    from docx.table import _Cell
    from docx.text.paragraph import Paragraph as BackendParagraph

    try:
        document = backend.Document(io.BytesIO(snapshot.content))
    except OpcError as error:
        raise DocumentParseError(f"Cannot read DOCX package: {error}") from error
    budget = InspectionBudget(limits)
    if document.element.body is None:
        raise DocumentParseError("DOCX has no document body")

    def blocks(element, owner):
        result = []
        for index, child in enumerate(element, 1):
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                budget.fact()
                paragraph = BackendParagraph(child, owner)
                budget.text(paragraph.text)
                runs = []
                for item in paragraph.iter_inner_content():
                    for run in getattr(item, "runs", (item,)):
                        budget.fact()
                        runs.append(
                            TextRun(
                                run.text, run.font.size.pt if run.font.size else None
                            )
                        )
                result.append(
                    Paragraph(
                        index,
                        paragraph.text,
                        tuple(runs),
                        paragraph.style.name if paragraph.style else None,
                    )
                )
            elif tag == "tbl":
                budget.fact()
                table = BackendTable(child, owner)
                cells, origins = [], {}
                for row_index, row in enumerate(table._tbl.tr_lst, 1):
                    column = row.grid_before
                    for node in row.tc_lst:
                        budget.fact()
                        column += 1
                        span = node.grid_span
                        if node.vMerge == "continue":
                            origin = origins.get(column)
                            if origin is not None:
                                old = cells[origin]
                                cells[origin] = TableCell(
                                    old.row,
                                    old.column,
                                    old.row_span + 1,
                                    old.column_span,
                                    old.blocks,
                                )
                        else:
                            cell = _Cell(node, table)
                            origins[column] = len(cells)
                            cells.append(
                                TableCell(
                                    row_index, column, 1, span, blocks(node, cell)
                                )
                            )
                        column += span - 1
                result.append(Table(index, tuple(cells)))
            elif tag in {"sdt", "sdtContent", "customXml"}:
                result.extend(blocks(child, owner))
        return tuple(result)

    return DocxReport(
        snapshot, blocks(document.element.body, document), limits.text_bytes
    )
