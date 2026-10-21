"""python-pptx adapter; no scoring policy."""

import io

from . import _pptx_styles as styles
from ._limits import DocumentLimitError, InspectionBudget
from .models import (
    Paragraph,
    PptxReport,
    PptxShape,
    PptxSlide,
    Table,
    TableCell,
    TextRun,
)


def colors(values, budget):
    def counted():
        for color in values:
            budget.color()
            yield tuple(color)

    return tuple(counted())


def style_diagnostic(error):
    """Budget exhaustion is fatal even during best-effort style resolution."""
    if isinstance(error, DocumentLimitError):
        raise error
    return str(error)


def inspect(snapshot, backend, options, limits):
    from lxml.etree import XMLSyntaxError
    from pptx.exc import InvalidXmlError

    from . import DocumentParseError

    try:
        presentation = backend.Presentation(io.BytesIO(snapshot.content))
        budget = InspectionBudget(limits)
        budget.pages(len(presentation.slides))
        slides = tuple(
            read_slide(presentation, slide, page, options, budget)
            for page, slide in enumerate(presentation.slides, 1)
        )
    except (XMLSyntaxError, InvalidXmlError) as error:
        raise DocumentParseError(f"Invalid PPTX XML: {error}") from error
    return PptxReport(snapshot, slides, limits.text_bytes)


def read_slide(presentation, slide, number, options, budget):
    entries = styles.leaves(slide.shapes, budget=budget)
    resolver = styles.Styles(presentation, slide, options, budget)
    backgrounds = styles.Backgrounds(presentation, slide, entries, resolver, budget)
    title_error = background_error = None
    chosen, title_paragraphs = None, []
    page_colors, title_background, title_colors = [], [], []
    try:
        chosen, title_paragraphs = styles.title(slide, entries, resolver)
    except ValueError as error:
        title_error = style_diagnostic(error)
    try:
        page_colors = backgrounds.page
    except ValueError as error:
        background_error = style_diagnostic(error)
    if chosen is not None:
        try:
            title_background = backgrounds.region(chosen)
        except ValueError as error:
            background_error = style_diagnostic(error)
            try:
                fill, opaque = resolver.shape_fill(chosen.shape, styles.WHITE)
                if opaque:
                    title_background = fill
            except ValueError as error:
                style_diagnostic(error)
    if chosen is None:
        title_error = title_error or "no identifiable page title"
    else:
        try:
            for paragraph in title_paragraphs:
                for _, properties in styles.runs(paragraph):
                    title_colors.extend(
                        colors(
                            resolver.font_colors(
                                chosen.shape, paragraph, properties, title_background
                            ),
                            budget,
                        )
                    )
        except ValueError as error:
            title_error = style_diagnostic(error)

    def paragraphs(shape, items, backs):
        result = []
        for index, paragraph in enumerate(items, 1):
            if not paragraph.text.strip():
                continue
            budget.fact()
            budget.text(paragraph.text)
            runs = []
            for text, properties in styles.runs(paragraph):
                budget.fact()
                size, ink, problem = None, (), None
                try:
                    size = resolver.font_size(shape, paragraph, properties)
                    ink = colors(
                        resolver.font_colors(shape, paragraph, properties, backs),
                        budget,
                    )
                except ValueError as error:
                    problem = style_diagnostic(error)
                runs.append(TextRun(text, size, ink, problem))
            result.append(Paragraph(index, paragraph.text, tuple(runs)))
        return tuple(result)

    shapes = []
    for index, entry in enumerate(entries, 1):
        shape = entry.shape
        backs = []
        if styles.paragraphs(shape):
            try:
                backs = backgrounds.region(entry)
            except ValueError as error:
                style_diagnostic(error)
        text, table = (), None
        if shape.has_text_frame:
            text = paragraphs(shape, shape.text_frame.paragraphs, backs)
        elif shape.has_table:
            budget.fact()
            cells = []
            for row_index, row in enumerate(shape.table.rows, 1):
                for column_index, cell in enumerate(row.cells, 1):
                    if not cell.is_spanned:
                        budget.fact()
                        cells.append(
                            TableCell(
                                row_index,
                                column_index,
                                cell.span_height,
                                cell.span_width,
                                paragraphs(shape, cell.text_frame.paragraphs, backs),
                            )
                        )
            table = Table(index, tuple(cells))
        shapes.append(
            PptxShape(
                index, shape.shape_id, entry.bounds, text, table, entry.geometry_error
            )
        )
    return PptxSlide(
        number,
        tuple(shapes),
        chosen.shape.shape_id if chosen else None,
        colors(title_colors, budget),
        title_error,
        colors(page_colors, budget),
        colors(title_background, budget),
        background_error,
    )
