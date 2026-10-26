"""Execute copied Harbor code in a process that cannot import Psycheval."""

from __future__ import annotations

import asyncio
import gzip
import hashlib
import importlib
import importlib.abc
import io
import json
import logging
import os
import pkgutil
import runpy
import shlex
import sys
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

PACKAGE = "downstream._vendor.harbor_copy"


class BlockPsycheval(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path, target=None):
        if fullname.split(".")[0] == "psycheval":
            raise AssertionError(f"copied code imported original package: {fullname}")
        return None


def module(name: str):
    return importlib.import_module(f"{PACKAGE}.{name}")


def run_module(name: str, *args: str) -> None:
    with patch.object(sys, "argv", [f"{PACKAGE}.{name}", *args]):
        try:
            runpy.run_module(f"{PACKAGE}.{name}", run_name="__main__", alter_sys=True)
        except SystemExit as exc:
            assert exc.code == 0, exc.code


def imports(root: Path) -> None:
    class BlockOptionalRuntime(importlib.abc.MetaPathFinder):
        def find_spec(self, fullname, path, target=None):
            if fullname.split(".")[0] in {
                "workbuddy_bench",
                "pypdf",
                "pptx",
                "openpyxl",
                "docx",
            }:
                raise ModuleNotFoundError("optional runtime is unavailable")
            return None

    sys.meta_path.insert(0, BlockOptionalRuntime())
    package = importlib.import_module(PACKAGE)
    found = list(pkgutil.walk_packages(package.__path__, f"{PACKAGE}."))
    for info in found:
        importlib.import_module(info.name)
    assert any(info.name == f"{PACKAGE}.tasks" for info in found)
    assert (
        module("agent").ExternalHarnessAgent(logs_dir=root, command="fixture").version()
        == package.__version__
    )
    runtime = module("runtime_config")
    assert runtime.HARNESS_PROTOCOL_VERSION == 2
    assert runtime.PEVAL_CONFIG_ENV == "PEVAL_CONFIG"
    for name in ("one.pptx", "two.pptx"):
        (root / name).write_bytes(b"identical without parser dependencies")
    selection = module("documents").select_document([root], format="pptx")
    assert selection.selected == root / "one.pptx"
    assert len(selection.candidates) == 2


def documents(root: Path) -> None:
    import docx
    import openpyxl
    import pptx
    import pypdf

    workdir = root / "workdir"
    workdir.mkdir()
    pdf = pypdf.PdfWriter()
    pdf.add_blank_page(width=100, height=200)
    pdf.write(workdir / "input.pdf")
    pptx.Presentation().save(workdir / "input.pptx")
    openpyxl.Workbook().save(workdir / "input.xlsx")
    docx.Document().save(workdir / "input.docx")
    api, evidence = module("documents"), module("evidence")
    mirror = root / "mirror"
    mirror.mkdir()
    for kind in ("pdf", "pptx", "xlsx", "docx"):
        (mirror / f"input.{kind}").write_bytes((workdir / f"input.{kind}").read_bytes())
    with evidence.EvidenceWriter(root / "evidence") as writer:
        for kind in ("pdf", "pptx", "xlsx", "docx"):
            selection = api.select_document([workdir, mirror], format=kind)
            assert selection.selected is not None and len(selection.candidates) == 2
            source = writer.capture(f"source_{kind}", selection.selected)
            report = api.inspect_document(source)
            writer.add_text(kind, report.to_markdown(), source_id=f"source_{kind}")
    assert (root / "evidence/artifact_manifest.json").is_file()


def native_office(root: Path) -> None:
    from harbor.models.task.config import EnvironmentConfig
    from harbor.models.task.task import Task
    from harbor.models.trial.paths import TrialPaths

    async def scenario():
        task = root / "bundle/tasks/office-00"
        paths = TrialPaths(root / "native trial 中文")
        paths.mkdir()
        environment = module("workbuddy_environment").WorkBuddyHostEnvironment(
            environment_dir=task / "environment",
            environment_name="copied-office",
            workdir_root=root / "workspaces",
            session_id="copied-office",
            trial_paths=paths,
            task_env_config=EnvironmentConfig(workdir="/workspace"),
            logger=logging.getLogger("copied-office"),
            host_access={"filesystem": True, "process": True},
            mounts=[
                {
                    "type": "bind",
                    "source": str(paths.agent_dir),
                    "target": "/logs/agent",
                },
                {
                    "type": "bind",
                    "source": str(paths.verifier_dir),
                    "target": "/logs/verifier",
                },
            ],
        )
        await environment.start(force_build=False)
        try:

            async def forbidden_shell(*args, **kwargs):
                raise AssertionError("native verifier invoked a shell")

            environment.exec = forbidden_shell
            verifier = module("workbuddy_verifier").NativePythonVerifier(
                task=Task(task),
                trial_paths=paths,
                environment=environment,
            )
            result = await verifier.verify()
            assert result.rewards["reward"] == 1.0, result
            assert (paths.verifier_dir / "office-adaptation.json").is_file()
        finally:
            await environment.stop(delete=True)

    asyncio.run(scenario())


def write_runtime(root: Path, *, action: str = "run") -> Path:
    runtime = module("runtime_config")
    path = runtime.write_effective_runtime_config(
        root / "runtime.json",
        runtime.EffectiveRuntimeConfig(
            paths=runtime.RuntimePaths(
                workdir=str(root),
                tests=str(root / "tests"),
                agent_logs=str(root / "agent"),
                verifier_logs=str(root / "verifier"),
                artifacts=str(root / "artifacts"),
            ),
            harness=runtime.HarnessInvocation(action=action),
        ),
    )
    os.environ["PEVAL_CONFIG"] = str(path)
    return path


def trajectory(message: str) -> dict:
    return {
        "schema_version": "ATIF-v1.7",
        "session_id": "copied-session",
        "agent": {"name": "downstream", "version": "1"},
        "steps": [{"step_id": 1, "source": "agent", "message": message}],
        "final_metrics": {"total_prompt_tokens": 12, "total_completion_tokens": 3},
    }


def trajectory_file(root: Path) -> None:
    assert sys.flags.utf8_mode == 0
    load = module("trajectory_validation").load_validated_trajectory
    directory = root / "trajectory"
    directory.mkdir()
    image = directory / "截图.png"
    image.write_bytes(b"fixture")
    payload = trajectory("完成 🎉")
    payload["steps"][0]["message"] = [
        {"type": "text", "text": "完成 🎉"},
        {"type": "image", "source": {"media_type": "image/png", "path": image.name}},
    ]
    path = directory / "trajectory.json"
    original = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    path.write_bytes(original)
    assert load(path) == payload
    image.unlink()
    try:
        load(path)
    except ValueError as exc:
        assert "referenced image file does not exist" in str(exc)
    else:
        raise AssertionError("copied validator accepted a missing relative image")
    assert path.read_bytes() == original


def verifier(root: Path) -> None:
    write_runtime(root)
    (root / "agent").mkdir()
    (root / "agent/trajectory.json").write_text(
        json.dumps(trajectory("done")), encoding="utf-8"
    )
    grader = root / "grader.json"
    for term, reward in (("done", 1), ("missing", 0)):
        grader.write_text(json.dumps({"final_terms": [term]}), encoding="utf-8")
        run_module("verifier", str(grader))
        assert (
            json.loads((root / "verifier/reward.json").read_text(encoding="utf-8"))[
                "reward"
            ]
            == reward
        )
        checks = json.loads((root / "verifier/score.json").read_text(encoding="utf-8"))[
            "verdicts"
        ]
        assert all(check["status"] == "pass" for check in checks) == bool(reward)
    grader.write_text(json.dumps({"custom_checks": ["value", "missing"]}))
    (root / "test_outputs.py").write_text(
        "import argparse, json; from pathlib import Path\n"
        "parser = argparse.ArgumentParser()\n"
        "parser.add_argument('--context', type=Path, required=True)\n"
        "parser.add_argument('--output', type=Path, required=True)\n"
        "args = parser.parse_args()\n"
        "context = json.loads(args.context.read_text(encoding='utf-8'))\n"
        "assert Path(context['paths']['agent_logs'], 'trajectory.json').is_file()\n"
        "args.output.write_text(json.dumps({'checks':[{'id':'value','passed':True}]}), encoding='utf-8')\n"
    )
    run_module("verifier", str(grader))
    assert (
        json.loads((root / "verifier/reward.json").read_text(encoding="utf-8"))[
            "reward"
        ]
        == 0.5
    )
    import yaml

    (root / "judge.yaml").write_text(
        yaml.safe_dump(
            {
                "version": 1,
                "artifacts": [
                    {
                        "id": "report",
                        "path": "missing.txt",
                        "type": "txt",
                        "required": True,
                    }
                ],
                "llm_judge": {
                    "method": "rubric_binary_mean",
                    "rubrics": [
                        {
                            "id": "quality",
                            "question": "Is the report supported?",
                            "artifact_refs": ["report"],
                            "pass_criteria": ["Supported facts"],
                            "fail_criteria": ["Missing facts"],
                        }
                    ],
                },
            }
        ),
        encoding="utf-8",
    )
    with patch.dict(
        os.environ,
        {
            "PEVAL_JUDGE_ENABLED": "true",
            "PEVAL_JUDGE_MODEL": "fixture",
            "PEVAL_JUDGE_BASE_URL": "http://127.0.0.1:1/v1",
        },
    ):
        # Missing required Agent evidence completes as failure without networking.
        run_module("verifier", str(grader))
    scores = json.loads((root / "verifier/reward.json").read_text(encoding="utf-8"))
    assert scores["rule_score"] == 0.5 and scores["llm_score"] == 0
    assert scores["reward"] == 0.4

    api = module("verifier")
    context = module("runtime_config").load_effective_runtime_config()

    async def explicit_configuration(use_manifest, custom_prompts=True):
        (root / "verifier/missing.txt").write_text(
            "copied evidence 中文 📰 {context}", encoding="utf-8"
        )
        if use_manifest:
            declaration = yaml.safe_load(
                (root / "judge.yaml").read_text(encoding="utf-8")
            )
            declaration["artifacts"][0].update(path="output/report.xlsx", type="xlsx")
            (root / "judge.yaml").write_text(yaml.safe_dump(declaration))
            script = root / "test_outputs.py"
            with script.open("a", encoding="utf-8") as stream:
                stream.write(
                    "logs = Path(context['paths']['verifier_logs'])\n"
                    "(logs / 'artifact_text').mkdir()\n"
                    "(logs / 'artifact_text/report.md').write_text('copied evidence 中文 📰 {context}', encoding='utf-8')\n"
                    "(logs / 'artifact_manifest.json').write_text(json.dumps({'artifacts':[{'id':'report','text_path':'artifact_text/report.md','extract_status':'ok'}]}))\n"
                )

        sent_messages = []

        async def request(client, settings, messages):
            sent_messages.append(messages)
            assert settings.reasoning_effort == "high"
            if custom_prompts:
                assert messages[0] == {"role": "system", "content": "copied system"}
                assert messages[1]["content"].startswith("copied\n")
                assert messages[1]["content"].endswith("\nend")
                supplied = json.loads(
                    messages[1]["content"][len("copied\n") : -len("\nend")]
                )
                assert (
                    supplied["evidence"]["report"]["text"]
                    == "copied evidence 中文 📰 {context}"
                )
            else:
                assert [message["role"] for message in messages] == ["user"]
                prompt = messages[0]["content"]
                assert prompt.startswith("You are a LLM-as-Judge evaluator.")
                assert (
                    "Faithful text:\n<artifact_faithful_text>\n"
                    "copied evidence 中文 📰 {context}\n</artifact_faithful_text>\n"
                ) in prompt
                rewards = json.loads(
                    prompt.split("## Rule reward\n", 1)[1].split(
                        "\n\n## Artifact evidence", 1
                    )[0]
                )
                assert rewards["reward"] == 0.5 and "llm_score" not in rewards
            return 200, json.dumps(
                {
                    "choices": [
                        {
                            "finish_reason": "stop",
                            "message": {
                                "content": json.dumps(
                                    {
                                        "id": "quality",
                                        "passed": True,
                                        "reason": "中文判断 📰",
                                        "passed_reason_note": None,
                                    }
                                )
                            },
                        }
                    ]
                }
            )

        with patch.object(module("verifier._judge"), "_request", request):
            result = await api.run_verifier(
                str(grader),
                runtime=context,
                judge_settings=api.JudgeSettings(
                    enabled=True,
                    base_url="http://127.0.0.1:1/v1",
                    model="copied-fixture",
                    reasoning_effort="high",
                    system_prompt="copied system" if custom_prompts else None,
                    user_prompt_template="copied\n{context}\nend"
                    if custom_prompts
                    else None,
                ),
            )
        assert json.loads(
            (root / "verifier/reward.json").read_text(encoding="utf-8")
        ) == {
            **scores,
            "llm_score": 1.0,
            "reward": 0.8 * 0.5 + 0.2,
        }
        (verdict,) = result["judges"][1]["verdicts"]
        assert verdict["item_id"] == "llm:quality"
        assert verdict["status"] == "pass" and verdict["score"] == 1
        assert verdict["reason"] == "中文判断 📰"
        raw = (root / "verifier/score.json").read_text(encoding="utf-8")
        assert "中文判断 📰" in raw
        assert result == json.loads(raw)
        index = json.loads(
            (root / "verifier/judge/evidence.json").read_text(encoding="utf-8")
        )["report"]
        relative = "artifact_text/report.md" if use_manifest else "missing.txt"
        raw = (root / "verifier" / relative).read_bytes()
        assert index["text_path"] == relative
        assert index["byte_count"] == len(raw)
        assert index["sha256"] == hashlib.sha256(raw).hexdigest()
        assert "text" not in index
        diagnostic = json.loads(
            (root / "verifier/judge/rubric-0001.json").read_text(encoding="utf-8")
        )
        assert diagnostic["messages"] == sent_messages[0]

    with patch.dict(
        os.environ,
        {"PEVAL_CONFIG": str(root / "missing.json"), "PEVAL_JUDGE_ENABLED": "invalid"},
    ):
        asyncio.run(explicit_configuration(False))
        asyncio.run(explicit_configuration(False, custom_prompts=False))
        asyncio.run(explicit_configuration(True))


def workbuddy(root: Path) -> None:
    from harbor.models.job.config import JobConfig

    work = module("workbuddy")
    bundle = root / "bundle"
    resolved = module("datasets").validate_harbor_dataset(
        dataset_id="fixture", path=bundle, format="workbuddy.v1"
    )
    output = root / "output"
    output.mkdir()
    config = output / "peval.toml"
    config.write_bytes(b"unreadable workspace config\xff")
    os.environ["PEVAL_CONFIG"] = str(config)
    base = JobConfig(
        datasets=[{"path": resolved.task_root}],
        agents=[{"name": "oracle"}],
        environment={
            "import_path": f"{PACKAGE}.environment:HostEnvironment",
            "kwargs": {"host_access": {"filesystem": True, "process": True}},
        },
    )
    before = base.model_dump()
    captured = io.StringIO()
    with redirect_stdout(captured):
        prepared = work.prepare_workbuddy_job(base)
        assert (
            prepared.environment.import_path
            == f"{PACKAGE}.workbuddy_environment:WorkBuddyHostEnvironment"
        )
        assert (
            prepared.verifier.import_path
            == f"{PACKAGE}.workbuddy_verifier:WorkBuddyVerifier"
        )
        assert prepared.n_attempts == 1 and prepared.timeout_multiplier == 1.0
        assert prepared.datasets == base.datasets and base.model_dump() == before
        compute = SimpleNamespace(
            compute_job_metrics=lambda path, expected_tasks: {
                "reward": 0,
                "n_tasks": len(expected_tasks),
            }
        )
        with patch.dict(sys.modules, {"workbuddy_bench.scorer.metrics": compute}):
            metrics = work.compute_official_metrics(
                output, expected_tasks=list(resolved.task_names)
            )
        assert metrics["n_tasks"] == len(resolved.task_names)
    assert captured.getvalue() == ""
    assert config.read_bytes() == b"unreadable workspace config\xff"
    assert {p.name for p in output.iterdir()} == {"peval.toml"}
    usage = module("workbuddy_usage")
    history = root / "history/project/session.jsonl"
    history.parent.mkdir(parents=True)
    history.write_text(
        json.dumps(
            {"providerData": {"messageId": "old", "rawUsage": {"prompt_tokens": 999}}}
        )
        + "\n",
        encoding="utf-8",
    )
    logs = root / "usage-logs"
    logs.mkdir()
    summary = {
        "invocation_id": "copy",
        "session_id": "session",
        "max_evidence_bytes": 1024 * 1024,
    }
    baseline = usage.history_baseline(root / "history", summary["max_evidence_bytes"])
    with history.open("a", encoding="utf-8") as stream:
        stream.write(
            json.dumps(
                {
                    "providerData": {
                        "messageId": "r",
                        "conversationRequestId": "turn",
                        "rawUsage": {
                            "prompt_tokens": 10,
                            "completion_tokens": 2,
                            "cached_tokens": 4,
                        },
                    }
                }
            )
            + "\n"
        )
    (logs / "acp.events.ndjson.gz").write_text(
        json.dumps(
            {
                "direction": "receive",
                "phase": "prompt",
                "message": {
                    "method": "session/update",
                    "params": {
                        "sessionId": "session",
                        "update": {
                            "sessionUpdate": "agent_message_chunk",
                            "_meta": {
                                "codebuddy.ai/modelRequestId": "r",
                                "codebuddy.ai/conversationRequestId": "turn",
                            },
                        },
                    },
                },
            }
        )
        + "\n",
        encoding="utf-8",
    )
    events = logs / "acp.events.ndjson.gz"
    events.write_bytes(gzip.compress(events.read_bytes(), mtime=0))
    evidence = usage.capture_usage(root / "history", logs, summary, baseline, "fixture")
    assert [record["model_request_id"] for record in evidence["records"]] == ["r"]
    digest = usage.publish_usage(logs / usage.USAGE_FILENAME, evidence)
    projected = usage.project_usage(
        {
            "steps": [
                {
                    "step_id": 1,
                    "source": "agent",
                    "message": "copy",
                    "extra": {"workbuddy_model_request_ids": ["r"]},
                }
            ],
            "extra": {},
        },
        evidence,
        digest,
    )
    assert projected["final_metrics"]["total_cached_tokens"] == 4
    assert (
        projected["steps"][0]["metrics"]["extra"]["model_inference"][
            "cache_read_tokens"
        ]
        == 4
    )


def synthetic_harness(root: Path) -> None:
    runtime = module("runtime_config").load_effective_runtime_config(
        require_harness=True
    )
    logs = Path(runtime.paths.agent_logs)
    state = logs / "state.json"
    previous = (
        json.loads(state.read_text()) if runtime.harness.action == "resume" else 0
    )
    state.write_text(json.dumps(previous + 1), encoding="utf-8")
    (logs / "trajectory.json").write_text(
        json.dumps(trajectory(sys.stdin.read())), encoding="utf-8"
    )


def host(root: Path) -> None:
    from harbor.models.agent.context import AgentContext
    from harbor.models.task.config import EnvironmentConfig
    from harbor.models.trial.paths import TrialPaths

    async def scenario() -> None:
        environment_dir = root / "environment"
        environment_dir.mkdir()
        (environment_dir / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
        paths = TrialPaths(root / "trial with spaces")
        paths.mkdir()
        artifacts = paths.artifacts_dir / "logs/artifacts"
        artifacts.mkdir(parents=True)
        environment = module("environment").HostEnvironment(
            environment_dir=environment_dir,
            environment_name="vendored",
            workdir_root=root / "workspaces",
            session_id="vendored",
            trial_paths=paths,
            task_env_config=EnvironmentConfig(workdir="/app"),
            logger=logging.getLogger("vendored"),
            host_access={"filesystem": True, "process": True},
            mounts=[
                {"type": "bind", "source": str(source), "target": target}
                for source, target in (
                    (paths.agent_dir, "/logs/agent"),
                    (paths.verifier_dir, "/logs/verifier"),
                    (artifacts, "/logs/artifacts"),
                )
            ],
        )
        agent = module("agent").ExternalHarnessAgent(
            logs_dir=paths.agent_dir,
            command=shlex.join(
                [
                    sys.executable,
                    "-I",
                    str(Path(__file__).resolve()),
                    str(root),
                    "synthetic_harness",
                ]
            ),
        )
        await environment.start(force_build=False)
        try:
            for action, instruction in (("run", "start"), ("resume", "continue")):
                context = AgentContext()
                await getattr(agent, action)(instruction, environment, context)
                assert context.n_input_tokens == 12 and context.n_output_tokens == 3
                assert context.metadata["harness_action"] == action
                result = json.loads((paths.agent_dir / "trajectory.json").read_text())
                assert result["steps"][0]["message"] == instruction
            assert json.loads((paths.agent_dir / "state.json").read_text()) == 2
        finally:
            await environment.stop(delete=True)

    asyncio.run(scenario())


def host_filesystem(root: Path) -> None:
    from harbor.models.task.config import EnvironmentConfig
    from harbor.models.trial.paths import TrialPaths

    async def scenario() -> None:
        environment_dir = root / "filesystem-environment"
        environment_dir.mkdir()
        (environment_dir / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
        paths = TrialPaths(root / "filesystem trial")
        paths.mkdir()
        artifacts = paths.artifacts_dir / "logs/artifacts"
        artifacts.mkdir(parents=True)
        environment = module("environment").HostEnvironment(
            environment_dir=environment_dir,
            environment_name="vendored-filesystem",
            workdir_root=root / "workspaces",
            session_id="vendored-filesystem",
            trial_paths=paths,
            task_env_config=EnvironmentConfig(workdir="/app"),
            logger=logging.getLogger("vendored-filesystem"),
            host_access={"filesystem": True, "process": False},
            workspace_baseline="none",
            mounts=[
                {"type": "bind", "source": str(source), "target": target}
                for source, target in (
                    (paths.agent_dir, "/logs/agent"),
                    (paths.verifier_dir, "/logs/verifier"),
                    (artifacts, "/logs/artifacts"),
                )
            ],
        )
        await environment.start(force_build=False)
        try:
            await environment.ensure_dirs(["/app/input"], chmod=False)
            environment.native_path("/app/input/value.txt").write_text(
                "value", encoding="utf-8"
            )
            assert await environment.is_file("/app/input/value.txt")
            downloaded = root / "downloaded.txt"
            await environment.download_file("/app/input/value.txt", downloaded)
            assert downloaded.read_text(encoding="utf-8") == "value"
            try:
                await environment.exec("true")
            except module("environment").HostProcessAccessError:
                pass
            else:
                raise AssertionError("filesystem-only Host executed a process")
        finally:
            await environment.stop(delete=True)

    asyncio.run(scenario())


def project(root: Path) -> None:
    from harbor.models.task.config import EnvironmentConfig
    from harbor.models.trial.paths import TrialPaths

    source = root / "project"
    source.mkdir()
    (source / "input.txt").write_text("original")
    task_context = root / "environment"
    task_context.mkdir()
    (task_context / "data").mkdir()
    (task_context / "data/task.txt").write_text("task input")
    (task_context / "prepare.py").write_text(
        "from pathlib import Path\ndef prepare(workdir):\n"
        "    source = Path(__file__).parent / 'data/task.txt'\n"
        "    (Path(workdir) / 'task.txt').write_text(source.read_text())\n"
    )
    paths = TrialPaths(root / "trial__YfQLWrD")
    paths.mkdir()
    os.environ["PEVAL_CONFIG"] = str(root / "missing.toml")
    environment = module("environment").HostEnvironment(
        environment_dir=task_context,
        environment_name="copied-project",
        session_id="project",
        trial_paths=paths,
        task_env_config=EnvironmentConfig(workdir="/workspace"),
        host_access={"filesystem": True, "process": True},
        workdir_root=root / "copies",
        workspace_source=source,
    )

    async def scenario():
        await environment.start(False)
        try:
            assert environment.work_dir == root / "copies" / "task_YfQLWrD"
            assert environment.work_dir == environment.native_path(".")
            result = await environment.exec_argv(
                [
                    sys.executable,
                    "-c",
                    "from pathlib import Path; assert Path('task.txt').read_text() == 'task input'; Path('input.txt').write_text('changed')",
                ]
            )
            assert result.return_code == 0, result.stderr
            assert (source / "input.txt").read_text() == "original"
        finally:
            await environment.stop(True)
        assert not list((root / "copies").iterdir())

    asyncio.run(scenario())


def opencode(root):
    from harbor.models.agent.context import AgentContext

    session = {
        "info": {"id": "copied-opencode", "outcome": "succeeded"},
        "messages": [
            {
                "id": "copied-message",
                "type": "assistant",
                "model": {"providerID": "fixture", "id": "model"},
                "time": {"created": 1, "completed": 2},
                "content": [{"type": "text", "text": "copied 完成"}],
                "tokens": {
                    "input": 10,
                    "output": 2,
                    "cache": {"read": 4, "write": 6},
                },
            }
        ],
    }
    (root / "opencode-session.json").write_text(json.dumps(session), encoding="utf-8")
    agent = module("opencode").HostOpenCodeAgent(
        logs_dir=root, model_name="fixture/model"
    )
    assert agent.options.websearch_provider == "random"
    assert module("opencode").HostOpenCodeOptions().websearch_provider == "random"
    assert (
        type(agent).parse_options({"websearch_provider": None}).websearch_provider
        is None
    )
    context = AgentContext()
    agent.populate_context_post_run(context)
    assert context.n_input_tokens == 20
    assert context.n_output_tokens == 2
    assert context.n_cache_tokens == 4
    data = module("trajectory_validation").load_validated_trajectory(
        root / "trajectory.json"
    )
    inference = data["final_metrics"]["extra"]["model_inference"]
    assert inference["cache_prompt_tokens"] == 20
    assert inference["cache_read_tokens"] == 4
    assert inference["cache_sample_count"] == 1


def acp(root):
    from harbor.models.agent.context import AgentContext
    from harbor.models.task.config import EnvironmentConfig
    from harbor.models.trial.paths import TrialPaths

    paths = TrialPaths(root / "trial")
    paths.mkdir()
    environment_dir = root / "environment"
    environment_dir.mkdir()
    (environment_dir / "Dockerfile").write_text("FROM scratch\n")
    # Emulate an installed downstream package for the child interpreter.
    host = module("environment").HostEnvironment(
        environment_dir=environment_dir,
        environment_name="copied-acp",
        session_id="copied",
        trial_paths=paths,
        task_env_config=EnvironmentConfig(workdir="/workspace"),
        host_access={"filesystem": True, "process": True},
        workdir_root=root / "work",
        workspace_baseline="none",
        persistent_env={"PYTHONPATH": str(root)},
    )
    received = []

    def observe(update):
        received.append(update)
        if update["sessionUpdate"] == "tool_call":
            (root / "live-update-ack").touch()

    agent = module("acp_agent").HostACPAgent(
        logs_dir=paths.agent_dir,
        argv=[
            sys.executable,
            str(root / "acp_peer.py"),
            "live",
            str(root / "live-update-ack"),
        ],
        on_session_update=observe,
    )

    async def scenario():
        await host.start(False)
        try:
            await agent.setup(host)
            for action in (agent.run, agent.resume):
                received.clear()
                context = AgentContext()
                await action("copied ACP", host, context)
                assert [update["sessionUpdate"] for update in received] == [
                    "agent_message_chunk",
                    "tool_call",
                    "tool_call_update",
                    "agent_message_chunk",
                ]
                assert received[1]["name"] == "bash"
                agent.populate_context_post_run(context)
                assert context.metadata["acp_stop_reason"] == "end_turn"
        finally:
            await host.stop(True)

    asyncio.run(scenario())


if __name__ == "__main__":
    sys.dont_write_bytecode = True
    root = Path(sys.argv[1])
    sys.path.insert(0, str(root))
    sys.meta_path.insert(0, BlockPsycheval())
    scenarios = {
        "imports": imports,
        "documents": documents,
        "trajectory_file": trajectory_file,
        "verifier": verifier,
        "workbuddy": workbuddy,
        "host": host,
        "host_filesystem": host_filesystem,
        "synthetic_harness": synthetic_harness,
        "native_office": native_office,
        "project": project,
        "opencode": opencode,
        "acp": acp,
    }
    scenarios[sys.argv[2]](root)
    assert not any(name.split(".")[0] == "psycheval" for name in sys.modules)
