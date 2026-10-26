"""Public document contracts, exercised with small synthetic documents."""

import dataclasses
import io
import os
import shutil
import subprocess
import sys
import zipfile
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace

import pytest

from psycheval.harbor.documents import (
    DocumentLimitError,
    DocumentLimits,
    DocumentParseError,
    MissingDocumentDependency,
    PptxOptions,
    inspect_document,
    select_document,
)
from psycheval.harbor.documents.models import Paragraph, Table
from psycheval.harbor.evidence import SourceSnapshot


@pytest.fixture(params=["pdf", "pptx", "xlsx", "docx"])
def limited_document(request):
    kind = request.param
    buffer = io.BytesIO()
    if kind == "pdf":
        pdf = pytest.importorskip("pypdf")
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        document = pdf.PdfWriter()
        for _ in range(2):
            page = document.add_blank_page(width=200, height=300)
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf (Enough text) Tj ET")
            page[NameObject("/Contents")] = stream
        document.write(buffer)
    elif kind == "pptx":
        pptx = pytest.importorskip("pptx")
        document = pptx.Presentation()
        for _ in range(2):
            slide = document.slides.add_slide(document.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 100, 100).text = "Enough text"
        document.save(buffer)
    elif kind == "xlsx":
        xlsx = pytest.importorskip("openpyxl")
        document = xlsx.Workbook()
        document.active["A1"] = "Enough text"
        document.create_sheet("Second")["A1"] = "Enough text"
        document.save(buffer)
    else:
        docx = pytest.importorskip("docx")
        document = docx.Document()
        document.add_paragraph("Enough text")
        document.add_table(1, 1).cell(0, 0).text = "More text"
        document.save(buffer)
    return SourceSnapshot(Path(f"limited.{kind}"), buffer.getvalue())


def test_document_fact_text_and_rendered_budgets(limited_document):
    for limits, message in (
        (DocumentLimits(facts=1), "facts"),
        (DocumentLimits(text_bytes=3), "text"),
    ):
        with pytest.raises(DocumentLimitError, match=message):
            inspect_document(limited_document, limits=limits)
    if limited_document.path.suffix != ".docx":
        with pytest.raises(DocumentLimitError, match="pages"):
            inspect_document(limited_document, limits=DocumentLimits(pages=1))
    report = inspect_document(limited_document)
    text = report.to_markdown()
    size = len(text.encode("utf-8"))
    assert dataclasses.replace(report, text_limit=size).to_markdown() == text
    with pytest.raises(DocumentLimitError, match="Markdown"):
        dataclasses.replace(report, text_limit=size - 1).to_markdown()


@pytest.mark.parametrize("snapshot", [False, True])
def test_inspection_source_limit_has_same_error_for_paths_and_snapshots(
    tmp_path, snapshot
):
    path = tmp_path / "source.pdf"
    path.write_bytes(b"too large")
    value = SourceSnapshot.read(path) if snapshot else path
    with pytest.raises(DocumentLimitError, match="exceeds"):
        inspect_document(value, limits=DocumentLimits(source_bytes=1))


def test_transitive_missing_dependency_names_the_broken_module(monkeypatch):
    import psycheval.harbor.documents as documents

    def missing(_):
        raise ModuleNotFoundError("missing backend dependency", name="PIL")

    monkeypatch.setattr(documents.importlib, "import_module", missing)
    with pytest.raises(MissingDocumentDependency, match="PIL"):
        inspect_document(SourceSnapshot(Path("file.pptx"), b"broken"))


def test_missing_docx_body_is_a_document_error():
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.element.remove(document.element.body)
    buffer = io.BytesIO()
    document.save(buffer)
    with pytest.raises(DocumentParseError, match="no document body"):
        inspect_document(SourceSnapshot(Path("bodyless.docx"), buffer.getvalue()))


def test_ooxml_package_with_no_workbook_is_a_document_error():
    docx = pytest.importorskip("docx")
    pytest.importorskip("openpyxl")
    buffer = io.BytesIO()
    docx.Document().save(buffer)
    with pytest.raises(DocumentParseError, match="workbook part"):
        inspect_document(SourceSnapshot(Path("renamed.xlsx"), buffer.getvalue()))


def test_source_permission_error_remains_an_infrastructure_error(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    source.touch()

    def denied(*args, **kwargs):
        raise PermissionError("fixture unreadable source")

    monkeypatch.setattr(Path, "open", denied)
    with pytest.raises(PermissionError, match="unreadable source"):
        inspect_document(source)


def test_unexpected_adapter_bug_is_not_scored_as_bad_output(monkeypatch):
    pdf = pytest.importorskip("pypdf")
    from psycheval.harbor.documents import _formats

    def broken(*_):
        raise TypeError("programming error")

    buffer = io.BytesIO()
    pdf.PdfWriter().write(buffer)
    monkeypatch.setattr(_formats, "inspect_pdf", broken)
    with pytest.raises(TypeError, match="programming error"):
        inspect_document(SourceSnapshot(Path("valid.pdf"), buffer.getvalue()))


@pytest.mark.parametrize("dependency", [False, True])
def test_pdf_backend_errors_distinguish_document_and_dependency(
    monkeypatch, dependency
):
    pdf = pytest.importorskip("pypdf")
    from pypdf.errors import DependencyError, PyPdfError

    def broken(*args, **kwargs):
        raise (DependencyError if dependency else PyPdfError)("backend diagnostic")

    monkeypatch.setattr(pdf, "PdfReader", broken)
    error = MissingDocumentDependency if dependency else DocumentParseError
    with pytest.raises(error, match="backend diagnostic"):
        inspect_document(SourceSnapshot(Path("file.pdf"), b"%PDF-1.7"))


@pytest.mark.skipif(sys.platform != "win32", reason="native Windows junction")
def test_selection_does_not_follow_native_junctions(tmp_path):
    root, outside = tmp_path / "workdir", tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    candidate = root / "one.pptx"
    candidate.touch()
    (outside / "two.pptx").touch()
    link = root / "linked"
    subprocess.run(
        ["cmd", "/c", "mklink", "/J", str(link), str(outside)],
        capture_output=True,
        check=True,
    )
    try:
        selection = select_document([root, link], format="pptx")
        assert selection.selected == candidate
        assert selection.candidates == (candidate,)
        assert (outside / "two.pptx").exists()
    finally:
        if link.is_junction():
            link.rmdir()


def test_selection_budgets_fail_closed_and_do_not_resolve_entries(
    tmp_path, monkeypatch
):
    (tmp_path / "one.pptx").touch()
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "two.pptx").touch()
    monkeypatch.setattr(
        Path, "resolve", lambda *a, **k: pytest.fail("unneeded realpath")
    )
    assert select_document([tmp_path], format="pptx", max_entries=1).selected is None
    assert "limit" in select_document([tmp_path], format="pptx", max_entries=1).reason
    limited = select_document([tmp_path], format="pptx", max_depth=0)
    assert limited.selected is None and "depth" in limited.reason
    assert len(select_document([tmp_path, nested], format="pptx").candidates) == 2


def test_unreadable_scan_cannot_claim_unique_candidate(tmp_path, monkeypatch):
    import psycheval.harbor.documents as documents

    candidate = tmp_path / "one.pptx"
    candidate.touch()
    blocked = tmp_path / "blocked"
    blocked.mkdir()
    scandir = documents.os.scandir

    def denied(path):
        if Path(path) == blocked:
            raise PermissionError("fixture denial")
        return scandir(path)

    monkeypatch.setattr(documents.os, "scandir", denied)
    result = select_document([candidate, blocked], format="pptx")
    assert result.selected is None and "PermissionError" in result.reason


def _pptx_snapshot(presentation):
    buffer = io.BytesIO()
    presentation.save(buffer)
    return SourceSnapshot(Path("synthetic.pptx"), buffer.getvalue())


@pytest.mark.parametrize("owner", ["master", "layout"])
@pytest.mark.parametrize("grouped", [False, True])
@pytest.mark.parametrize("pages", [1, 2])
def test_pptx_inherited_shapes_share_the_document_facts_budget(owner, grouped, pages):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[6]
    slides = [presentation.slides.add_slide(layout) for _ in range(pages)]
    # Keep theme/text styles but remove template placeholders for an exact budget.
    for collection in (layout.shapes, layout.slide_master.shapes):
        for shape in list(collection):
            collection._spTree.remove(shape._element)
    shapes = slides[0].shapes
    children = shapes.add_group_shape().shapes if grouped else shapes
    for index in range(12):
        children.add_shape(MSO_SHAPE.RECTANGLE, index * 10, 0, 100, 100)
    inherited = layout.slide_master.shapes if owner == "master" else layout.shapes
    for shape in list(shapes):
        inherited._spTree.append(shape._element)
    snapshot = _pptx_snapshot(presentation)
    facts = (12 + int(grouped)) * pages
    assert (
        len(inspect_document(snapshot, limits=DocumentLimits(facts=facts)).slides)
        == pages
    )
    with pytest.raises(DocumentLimitError, match="facts"):
        inspect_document(snapshot, limits=DocumentLimits(facts=facts - 1))

    # An inheritance tree disabled for background analysis contributes no work.
    for slide in slides:
        slide._element.set("showMasterSp", "0")
    assert (
        len(
            inspect_document(
                _pptx_snapshot(presentation), limits=DocumentLimits(facts=1)
            ).slides
        )
        == pages
    )


def test_utf8_budgets_count_multibyte_text_and_preserve_remaining_capacity():
    from psycheval.harbor.documents._limits import InspectionBudget, bounded_markdown

    budget = InspectionBudget(DocumentLimits(text_bytes=7))
    budget.text("中")  # 3 bytes
    with pytest.raises(DocumentLimitError, match="text"):
        budget.text("中文")  # Would need 6 more bytes.
    assert budget.text_bytes == 3
    budget.text("🙂")  # Exactly the remaining 4 bytes.
    assert budget.text_bytes == 7
    with pytest.raises(DocumentLimitError, match="text"):
        budget.text("a")
    assert bounded_markdown(["中", "🙂"], 9) == "中\n\n🙂"
    with pytest.raises(DocumentLimitError, match="Markdown"):
        bounded_markdown(["中", "🙂"], 8)


def test_obviously_oversized_text_is_rejected_without_encoding():
    from psycheval.harbor.documents._limits import InspectionBudget, bounded_markdown

    class NoEncoding(str):
        def __str__(self):
            return self

        def encode(self, *args, **kwargs):
            pytest.fail("oversized text must not allocate an encoded copy")

    text = NoEncoding("already too large")
    with pytest.raises(DocumentLimitError, match="text"):
        InspectionBudget(DocumentLimits(text_bytes=1)).text(text)
    with pytest.raises(DocumentLimitError, match="Markdown"):
        bounded_markdown([text], 1)


def test_pptx_translucent_palette_expansion_consumes_a_document_budget():
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement

    def document(overlays):
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for index in range(overlays):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                0,
                0,
                presentation.slide_width // 2,
                presentation.slide_height,
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(
                10 + index * 13, 100 + index * 5, 220 - index * 7
            )
            alpha = OxmlElement("a:alpha")
            alpha.set("val", "50000")
            shape.fill.fore_color._color._xClr.append(alpha)
        slide.shapes.add_textbox(
            0, 0, presentation.slide_width, presentation.slide_height
        ).text = "Heading"
        return _pptx_snapshot(presentation)

    small = inspect_document(document(2)).slides[0]
    assert set(small.title_background_colors) == {
        (255, 255, 255),
        (132.5, 177.5, 237.5),
        (139, 180, 234),
        (77.75, 141.25, 225.25),
    }
    large = document(17)
    with pytest.raises(DocumentLimitError, match="color operations"):
        inspect_document(large)
    expanded = inspect_document(large, limits=DocumentLimits(color_operations=1000000))
    assert len(expanded.slides[0].title_background_colors) > 10000


@pytest.mark.parametrize("owner", ["slide", "layout", "master"])
def test_pptx_color_budget_boundary_and_cross_slide_inheritance(owner):
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    target = {
        "slide": slide,
        "layout": slide.slide_layout,
        "master": slide.slide_layout.slide_master,
    }[owner]
    target.background.fill.solid()
    target.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    snapshot = _pptx_snapshot(presentation)
    # Resolution, composition, palette collection and report materialization.
    limits = DocumentLimits(color_operations=4)
    assert inspect_document(snapshot, limits=limits).slides[0].background_colors == (
        (255, 255, 255),
    )
    with pytest.raises(DocumentLimitError, match="color operations"):
        inspect_document(snapshot, limits=DocumentLimits(color_operations=3))
    second = presentation.slides.add_slide(slide.slide_layout)
    if owner == "slide":
        second.background.fill.solid()
        second.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    with pytest.raises(DocumentLimitError, match="color operations"):
        inspect_document(_pptx_snapshot(presentation), limits=limits)


def test_pptx_duplicate_gradient_stops_still_consume_color_work():
    pptx = pytest.importorskip("pptx")
    from copy import deepcopy

    from pptx.dml.color import RGBColor

    from psycheval.harbor.documents._pptx_styles import NS

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.gradient()
    for stop in slide.background.fill.gradient_stops:
        stop.color.rgb = RGBColor(255, 255, 255)
    stops = slide._element.find("p:cSld/p:bg/p:bgPr/a:gradFill/a:gsLst", NS)
    for _ in range(100):
        stops.append(deepcopy(stops[0]))
    snapshot = _pptx_snapshot(presentation)
    assert inspect_document(snapshot).slides[0].background_colors == ((255, 255, 255),)
    with pytest.raises(DocumentLimitError, match="color operations"):
        inspect_document(snapshot, limits=DocumentLimits(color_operations=10))


def test_pptx_palette_collection_consumes_candidates_incrementally():
    from psycheval.harbor.documents._limits import InspectionBudget
    from psycheval.harbor.documents._pptx_styles import distinct

    consumed = []

    def candidates():
        for index in range(100):
            consumed.append(index)
            yield (255, 255, 255)

    budget = InspectionBudget(DocumentLimits(color_operations=3))
    with pytest.raises(DocumentLimitError, match="color operations"):
        distinct(candidates(), budget)
    assert consumed == [0, 1, 2, 3]


@pytest.mark.parametrize("target", ["title", "body"])
def test_pptx_text_color_budget_failure_is_not_a_style_diagnostic(target):
    pptx = pytest.importorskip("pptx")
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("title", "body")):
        shape = slide.shapes.add_textbox(0, Inches(index), Inches(5), Inches(1))
        shape.text = name
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(32 if name == "title" else 18)
        if name == target:
            gradient = OxmlElement("a:gradFill")
            stops = OxmlElement("a:gsLst")
            gradient.append(stops)
            for index in range(100):
                stop = OxmlElement("a:gs")
                stop.set("pos", str(index * 1000))
                color = OxmlElement("a:srgbClr")
                color.set("val", "0000FF")
                stop.append(color)
                stops.append(stop)
            run._r.get_or_add_rPr().append(gradient)
    snapshot = _pptx_snapshot(presentation)
    report = inspect_document(snapshot)
    assert report.slides[0].title_error is None
    assert report.slides[0].shapes[1].paragraphs[0].runs[0].error is None
    with pytest.raises(DocumentLimitError, match="color operations"):
        inspect_document(snapshot, limits=DocumentLimits(color_operations=100))


def test_pptx_reuses_slide_backgrounds_and_keeps_distinct_colors(monkeypatch):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    from psycheval.harbor.documents import _pptx_styles

    presentation = pptx.Presentation()
    pages, shapes = 4, 14
    for _ in range(pages):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for index in range(shapes):
            shape = slide.shapes.add_textbox(
                Inches(index * 0.01), Inches(index * 0.01), Inches(5), Inches(1)
            )
            shape.text = f"Text {index}"
    calls = []
    original = _pptx_styles.background_shapes

    def counted(*args):
        calls.append(1)
        return original(*args)

    monkeypatch.setattr(_pptx_styles, "background_shapes", counted)
    report = inspect_document(_pptx_snapshot(presentation))
    assert len(calls) == pages
    assert all(
        len(run.colors) == 1
        for slide in report.slides
        for shape in slide.shapes
        for paragraph in shape.paragraphs
        for run in paragraph.runs
    )


@pytest.mark.parametrize("problem", ["chOff", "chExt", "zero_extent"])
def test_pptx_unknown_group_geometry_retains_body_text(problem):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    from psycheval.harbor.documents._pptx_styles import NS

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Readable group text"
    transform = group._element.find("p:grpSpPr/a:xfrm", NS)
    if problem == "zero_extent":
        transform.find("a:chExt", NS).set("cx", "0")
    else:
        transform.remove(transform.find(f"a:{problem}", NS))
    report = inspect_document(_pptx_snapshot(presentation))
    assert "Readable group text" in report.to_markdown()
    shape = report.slides[0].shapes[0]
    assert shape.bounds is None and shape.geometry_error
    assert report.slides[0].text_shape_count == 1
    assert report.slides[0].background_error


def test_pptx_bad_title_size_does_not_hide_page_background():
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(
        0, 0, Inches(2), Inches(1)
    ).text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Still readable"
    run._r.get_or_add_rPr().set("sz", "invalid")
    report = inspect_document(_pptx_snapshot(presentation))
    page = report.slides[0]
    assert page.title_error
    assert page.background_error is None
    assert page.background_colors == ((255, 255, 255),)
    assert "Still readable" in report.to_markdown()


@pytest.mark.parametrize("format", ["pdf", "pptx", "xlsx", "docx"])
def test_selection_is_unique_and_deduplicates_roots(tmp_path, format):
    nested = tmp_path / "nested"
    nested.mkdir()
    first = nested / f"first.{format.upper()}"
    first.write_bytes(b"first")
    (nested / f"~$lock.{format}").touch()
    assert select_document([tmp_path, nested], format=format).selected == first
    second = tmp_path / f"second.{format}"
    second.write_bytes(b"other")
    result = select_document([tmp_path], format=format)
    assert result.selected is None and len(result.candidates) == 2
    assert (
        select_document([tmp_path], format=format, preferred_stem="first").selected
        == first
    )
    (tmp_path / f"first.{format}").write_bytes(b"third")
    assert (
        select_document([tmp_path], format=format, preferred_stem="first").selected
        is None
    )


@pytest.mark.parametrize("format", ["pdf", "pptx", "xlsx", "docx"])
@pytest.mark.parametrize("copies", [2, 4])
@pytest.mark.parametrize("same_name", [False, True])
def test_selection_collapses_identical_workspace_and_mirror_copies(
    tmp_path, format, copies, same_name
):
    workdir, artifacts = tmp_path / "workdir", tmp_path / "artifacts"
    paths = []
    for index in range(copies):
        root = workdir if index % 2 == 0 else artifacts
        name = "report" if same_name else f"report-{index}"
        path = root / str(index // 2) / f"{name}.{format}"
        path.parent.mkdir(parents=True)
        path.write_bytes(b"identical document content")
        paths.append(path)
    expected = tuple(sorted(paths))
    for roots in ([workdir, artifacts], [artifacts, workdir]):
        result = select_document(roots, format=format)
        assert result.selected == expected[0]
        assert result.candidates == expected
        assert result.reason == "unique after content dedupe"


@pytest.mark.parametrize("format", ["pdf", "pptx", "xlsx", "docx"])
@pytest.mark.parametrize("copies", [1, 2])
def test_selection_preserves_preferred_stem_before_content_comparison(
    tmp_path, format, copies
):
    alias = tmp_path / "aaa" / f"alias.{format}"
    alias.parent.mkdir()
    alias.write_bytes(b"preferred document")
    (tmp_path / f"other.{format}").write_bytes(b"different document")
    preferred = []
    for index in range(copies):
        path = tmp_path / f"preferred-{index}" / f"final.{format}"
        path.parent.mkdir()
        path.write_bytes(alias.read_bytes())
        preferred.append(path)
    result = select_document([tmp_path], format=format, preferred_stem="final")
    assert result.candidates[0] == alias
    assert result.selected == preferred[0]
    assert len(result.candidates) == copies + 2


@pytest.mark.parametrize("other", [b"different size", b"other"])
def test_selection_rejects_different_contents(tmp_path, other):
    first, second = tmp_path / "one.pptx", tmp_path / "two.pptx"
    first.write_bytes(b"first")
    second.write_bytes(other)
    result = select_document([tmp_path], format="pptx")
    assert result.selected is None
    assert result.reason == "ambiguous PPTX candidates"


def test_selection_compares_contents_beyond_first_chunk(tmp_path):
    first, second = tmp_path / "one.pptx", tmp_path / "two.pptx"
    prefix = b"x" * (1024 * 1024 + 1)
    first.write_bytes(prefix + b"a")
    second.write_bytes(prefix + b"b")
    assert select_document([tmp_path], format="pptx").selected is None
    second.write_bytes(first.read_bytes())
    assert select_document([tmp_path], format="pptx").selected == first


@pytest.mark.parametrize(
    "format,module",
    [("pdf", "pypdf"), ("pptx", "pptx"), ("xlsx", "openpyxl"), ("docx", "docx")],
)
def test_identical_empty_documents_are_selected_but_fail_inspection(
    tmp_path, format, module
):
    pytest.importorskip(module)
    first, second = tmp_path / f"one.{format}", tmp_path / f"two.{format}"
    first.touch()
    second.touch()
    result = select_document([tmp_path], format=format, max_compare_bytes=1)
    assert result.selected == first
    with pytest.raises(DocumentParseError):
        inspect_document(result.selected)


def test_selection_comparison_budget_counts_repeated_reads(tmp_path):
    paths = [tmp_path / f"{index}.pptx" for index in range(3)]
    for path in paths:
        path.write_bytes(b"12345")
    result = select_document([tmp_path], format="pptx", max_compare_bytes=19)
    assert result.selected is None and result.candidates == tuple(paths)
    assert "incomplete document comparison" in result.reason
    assert "limit" in result.reason
    assert (
        select_document([tmp_path], format="pptx", max_compare_bytes=20).selected
        == paths[0]
    )


@pytest.mark.parametrize("limit", [0, -1, True, 1.5, None])
def test_selection_rejects_invalid_comparison_budget(tmp_path, limit):
    with pytest.raises(ValueError, match="max_compare_bytes"):
        select_document([tmp_path], format="pptx", max_compare_bytes=limit)


def test_selection_does_not_read_singletons_or_different_sizes(tmp_path, monkeypatch):
    import psycheval.harbor.documents as documents

    first, second = tmp_path / "one.pptx", tmp_path / "two.pptx"
    first.write_bytes(b"a")
    second.write_bytes(b"bb")
    monkeypatch.setattr(
        documents.os, "open", lambda *a, **k: pytest.fail("unnecessary content read")
    )
    assert select_document([first], format="pptx").selected == first
    assert select_document([tmp_path], format="pptx").selected is None
    assert (
        select_document([tmp_path], format="pptx", preferred_stem="one").selected
        == first
    )


@pytest.mark.parametrize("error", [PermissionError, FileNotFoundError])
def test_selection_cannot_ignore_unreadable_copy(tmp_path, monkeypatch, error):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.open

    def denied(path, *args, **kwargs):
        if Path(path) == paths[1]:
            raise error("fixture denial")
        return original(path, *args, **kwargs)

    monkeypatch.setattr(documents.os, "open", denied)
    result = select_document([tmp_path], format="pptx")
    assert result.selected is None and result.candidates == tuple(paths)
    assert "incomplete document comparison" in result.reason
    assert error.__name__ in result.reason


@pytest.mark.parametrize("replacement", [b"short", b"different"])
def test_selection_detects_changes_between_discovery_and_open(
    tmp_path, monkeypatch, replacement
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.open

    def change(path, *args, **kwargs):
        if Path(path) == paths[1]:
            paths[1].write_bytes(replacement)
            info = paths[1].stat()
            os.utime(paths[1], ns=(info.st_atime_ns, info.st_mtime_ns + 1_000_000_000))
        return original(path, *args, **kwargs)

    monkeypatch.setattr(documents.os, "open", change)
    result = select_document([tmp_path], format="pptx")
    assert result.selected is None
    assert "incomplete document comparison" in result.reason
    assert "changed" in result.reason


def test_selection_rechecks_earlier_copies_after_later_comparisons(
    tmp_path, monkeypatch
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(3)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.open

    def change(path, *args, **kwargs):
        if Path(path) == paths[2]:
            paths[1].write_bytes(b"changed after its comparison")
        return original(path, *args, **kwargs)

    monkeypatch.setattr(documents.os, "open", change)
    result = select_document([tmp_path], format="pptx")
    assert result.selected is None
    assert "incomplete document comparison" in result.reason
    assert "changed" in result.reason


@pytest.mark.parametrize("failure", ["write", "read_error"])
def test_selection_detects_failure_during_content_read(tmp_path, monkeypatch, failure):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.fdopen
    changed = False

    @contextmanager
    def wrap(*args, **kwargs):
        with original(*args, **kwargs) as stream:

            def read(count):
                nonlocal changed
                data = stream.read(count)
                if not changed:
                    changed = True
                    if failure == "read_error":
                        raise OSError("fixture read failure")
                    paths[0].write_bytes(b"different")
                    info = paths[0].stat()
                    os.utime(
                        paths[0],
                        ns=(info.st_atime_ns, info.st_mtime_ns + 1_000_000_000),
                    )
                return data

            yield SimpleNamespace(read=read, fileno=stream.fileno)

    monkeypatch.setattr(documents.os, "fdopen", wrap)
    result = select_document([tmp_path], format="pptx")
    assert result.selected is None and changed
    assert "incomplete document comparison" in result.reason
    assert (
        "changed" if failure == "write" else "fixture read failure"
    ) in result.reason


def test_selection_does_not_compare_path_ctime_with_descriptor_ctime(
    tmp_path, monkeypatch
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.fstat

    def descriptor_info(descriptor):
        info = original(descriptor)
        return SimpleNamespace(
            **{
                name: getattr(info, name)
                for name in dir(info)
                if name.startswith("st_") and name != "st_ctime_ns"
            },
            st_ctime_ns=info.st_ctime_ns + 1_000_000_000,
        )

    monkeypatch.setattr(documents.os, "fstat", descriptor_info)
    assert select_document(paths, format="pptx").selected == paths[0]


def test_selection_rejects_opening_another_file_with_identical_bytes_and_mtime(
    tmp_path, monkeypatch
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    replacement = tmp_path / "replacement"
    shutil.copy2(paths[1], replacement)
    assert replacement.stat().st_mtime_ns == paths[1].stat().st_mtime_ns
    assert not os.path.samestat(paths[1].stat(), replacement.stat())
    original = os.open

    def substitute(path, *args, **kwargs):
        # Model a path replacement at open, restored before later path checks.
        return original(
            replacement if Path(path) == paths[1] else path, *args, **kwargs
        )

    monkeypatch.setattr(documents.os, "open", substitute)
    result = select_document(paths, format="pptx")
    assert result.selected is None and result.candidates == tuple(paths)
    assert "incomplete document comparison" in result.reason
    assert "changed" in result.reason


def test_selection_rejects_descriptor_change_even_if_path_state_is_restored(
    tmp_path, monkeypatch
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.fstat
    seen = set()

    def changed_descriptor(descriptor):
        info = original(descriptor)
        if descriptor not in seen:
            seen.add(descriptor)
            return info
        return SimpleNamespace(
            **{
                name: getattr(info, name)
                for name in dir(info)
                if name.startswith("st_") and name != "st_ctime_ns"
            },
            st_ctime_ns=info.st_ctime_ns + 1_000_000_000,
        )

    monkeypatch.setattr(documents.os, "fstat", changed_descriptor)
    result = select_document(paths, format="pptx")
    assert result.selected is None
    assert "incomplete document comparison" in result.reason
    assert "changed" in result.reason


def test_selection_closes_descriptor_when_stream_creation_fails(tmp_path, monkeypatch):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    opened = []

    def fail(descriptor, *args, **kwargs):
        opened.append(descriptor)
        raise OSError("fixture fdopen failure")

    monkeypatch.setattr(documents.os, "fdopen", fail)
    try:
        result = select_document(paths, format="pptx")
        assert result.selected is None and "fixture fdopen failure" in result.reason
        assert len(opened) == 1
        with pytest.raises(OSError):
            os.fstat(opened[0])
    finally:
        for descriptor in opened:
            try:
                os.close(descriptor)
            except OSError:
                pass


def test_selection_reports_short_read_without_claiming_file_changed(
    tmp_path, monkeypatch
):
    import psycheval.harbor.documents as documents

    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.fdopen

    @contextmanager
    def short_read(*args, **kwargs):
        with original(*args, **kwargs) as stream:
            yield SimpleNamespace(
                read=lambda count: stream.read(count)[:1], fileno=stream.fileno
            )

    monkeypatch.setattr(documents.os, "fdopen", short_read)
    result = select_document(paths, format="pptx")
    assert result.selected is None
    assert "incomplete document comparison" in result.reason
    assert "short read" in result.reason


def test_selection_rejects_candidate_replaced_with_link(tmp_path, monkeypatch):
    import psycheval.harbor.documents as documents

    outside = tmp_path / "outside"
    outside.write_bytes(b"identical")
    link = tmp_path / "link"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symlink creation unavailable")
    link.unlink()
    paths = [tmp_path / f"{index}.pptx" for index in range(2)]
    for path in paths:
        path.write_bytes(b"identical")
    original = os.open

    def replace(path, *args, **kwargs):
        if Path(path) == paths[0]:
            paths[1].unlink()
            paths[1].symlink_to(outside)
        return original(path, *args, **kwargs)

    monkeypatch.setattr(documents.os, "open", replace)
    result = select_document(paths, format="pptx")
    assert result.selected is None and result.candidates == tuple(paths)
    assert "incomplete document comparison" in result.reason
    assert "Linked" in result.reason
    assert outside.read_bytes() == b"identical"


@pytest.mark.parametrize(
    "format,module",
    [("pdf", "pypdf"), ("pptx", "pptx"), ("xlsx", "openpyxl"), ("docx", "docx")],
)
def test_missing_dependency_is_not_a_parse_failure(monkeypatch, format, module):
    import psycheval.harbor.documents as documents

    original = documents.importlib.import_module

    def missing(name):
        if name == module:
            raise ModuleNotFoundError(name)
        return original(name)

    monkeypatch.setattr(documents.importlib, "import_module", missing)
    with pytest.raises(MissingDocumentDependency, match=rf"psycheval\[{format}\]"):
        inspect_document(SourceSnapshot(Path(f"file.{format}"), b"broken"))


@pytest.mark.parametrize(
    "format,module",
    [("pdf", "pypdf"), ("pptx", "pptx"), ("xlsx", "openpyxl"), ("docx", "docx")],
)
def test_corrupt_documents_have_a_common_error(format, module):
    pytest.importorskip(module)
    with pytest.raises(DocumentParseError):
        inspect_document(SourceSnapshot(Path(f"bad.{format}"), b"not a document"))


def test_pdf_empty_pages_dimensions_and_encryption(tmp_path):
    pdf = pytest.importorskip("pypdf")
    writer = pdf.PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.add_blank_page(width=400, height=500)
    path = tmp_path / "file.pdf"
    writer.write(path)
    report = inspect_document(path)
    assert [(p.number, p.text, p.width, p.height) for p in report.pages] == [
        (1, "", 200, 300),
        (2, "", 400, 500),
    ]
    assert "## Page 2" in report.to_markdown()
    with pytest.raises(dataclasses.FrozenInstanceError):
        report.pages[0].text = "mutated"
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(DocumentParseError, match="Encrypted"):
        inspect_document(path)


def test_xlsx_formulas_cached_values_merges_and_hidden_sheets(tmp_path):
    xlsx = pytest.importorskip("openpyxl")
    book = xlsx.Workbook()
    sheet = book.active
    sheet.title = "中文"
    sheet["A1"] = "Heading"
    sheet.merge_cells("A1:B1")
    sheet["A2"] = 12
    sheet["B2"] = "=A2*2"
    book.create_sheet("hidden").sheet_state = "hidden"
    path = tmp_path / "book.xlsx"
    book.save(path)
    report = inspect_document(path)
    assert [(s.name, s.visibility) for s in report.sheets] == [
        ("中文", "visible"),
        ("hidden", "hidden"),
    ]
    assert report.sheets[0].merged_ranges == ("A1:B1",)
    formula = report.sheets[0].cells[-1]
    assert (formula.address, formula.value, formula.formula, formula.cached_value) == (
        "B2",
        None,
        "=A2*2",
        None,
    )
    assert "Cached value: None" in report.to_markdown()


def test_sparse_xlsx_and_formula_cache_are_not_recalculated(tmp_path):
    xlsx = pytest.importorskip("openpyxl")
    book = xlsx.Workbook()
    sheet = book.active
    sheet["A1"] = "=12*2"
    sheet["XFD1048576"] = "Distant"
    buffer = io.BytesIO()
    book.save(buffer)
    result = io.BytesIO()
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(result, "w") as modified:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>12*2</f><v></v>", b"<f>12*2</f><v>999</v>")
            modified.writestr(name, data)
    report = inspect_document(
        SourceSnapshot(tmp_path / "sparse.xlsx", result.getvalue())
    )
    cells = report.sheets[0].cells
    assert len(cells) == 2 and cells[1].address == "XFD1048576"
    assert cells[0].formula == "=12*2" and cells[0].cached_value == 999


def test_docx_body_order_nested_tables_and_spans(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("Intro", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Vertical"
    nested = table.cell(0, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).merge(nested.cell(0, 1)).text = "Nested"
    document.add_paragraph("Last")
    document.sections[0].header.paragraphs[0].text = "HEADER_ONLY"
    path = tmp_path / "doc.docx"
    document.save(path)
    report = inspect_document(path)
    assert [type(block) for block in report.blocks] == [Paragraph, Table, Paragraph]
    assert report.blocks[0].style == "Heading 1"
    assert report.blocks[1].cells[0].row_span == 2
    assert any(
        isinstance(b, Table) and b.cells[0].column_span == 2
        for b in report.blocks[1].cells[1].blocks
    )
    text = report.to_markdown()
    assert (
        text.index("Intro")
        < text.index("Vertical")
        < text.index("Nested")
        < text.index("Last")
    )
    assert "HEADER_ONLY" not in text


def test_pptx_reports_body_title_and_independent_unknown_colors(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for size, text in [(32, "Title"), (18, "Body")]:
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 255)
    slide.notes_slide.notes_text_frame.text = "NOTES_ONLY"
    path = tmp_path / "deck.pptx"
    deck.save(path)
    report = inspect_document(path, pptx_options=PptxOptions(18, 0.95))
    page = report.slides[0]
    assert page.title_shape_id == slide.shapes[0].shape_id
    assert page.text_shape_count == 2 and page.title_colors == ((0, 0, 255),)
    assert page.background_colors == ((255, 255, 255),)
    assert page.shapes[0].paragraphs[0].runs[0].font_size == 32
    assert "NOTES_ONLY" not in report.to_markdown()
    path.write_bytes(b"changed")
    assert inspect_document(report.source).to_markdown() == report.to_markdown()


@pytest.mark.parametrize("decoration", ["01", "07.", "１２", "•", "—"])
def test_pptx_text_heading_precedes_larger_decorative_labels(tmp_path, decoration):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches, Pt

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for text, size in [(decoration, 50), ("LLM 与 Token", 29.2), ("正文", 18)]:
        shape = slide.shapes.add_textbox(0, 0, Inches(5), Inches(1))
        shape.text = text
        shape.text_frame.paragraphs[0].font.size = Pt(size)
    path = tmp_path / "title.pptx"
    deck.save(path)
    page = inspect_document(path).slides[0]
    assert page.title_shape_id == slide.shapes[1].shape_id
    assert page.text_shape_count == 3


@pytest.mark.parametrize("placeholder", [False, True])
def test_pptx_numeric_title_remains_identifiable(tmp_path, placeholder):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches, Pt

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1 if placeholder else 6])
    title = (
        slide.shapes.title
        if placeholder
        else slide.shapes.add_textbox(0, 0, Inches(5), Inches(1))
    )
    title.text = "2026"
    title.text_frame.paragraphs[0].font.size = Pt(32)
    body = slide.shapes.add_textbox(0, Inches(2), Inches(5), Inches(1))
    body.text = "Body" if placeholder else "01"
    body.text_frame.paragraphs[0].font.size = Pt(18)
    path = tmp_path / "numeric.pptx"
    deck.save(path)
    assert inspect_document(path).slides[0].title_shape_id == title.shape_id


def test_archive_limits_fail_without_partial_reports():
    pytest.importorskip("pptx")
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr("one", "expanded text")
        archive.writestr("two", "expanded text")
    source = SourceSnapshot(Path("deck.pptx"), content.getvalue())
    for limits in (
        DocumentLimits(archive_entries=1),
        DocumentLimits(expanded_bytes=1),
        DocumentLimits(source_bytes=1),
    ):
        with pytest.raises(DocumentParseError, match="limit"):
            inspect_document(source, limits=limits)


@pytest.mark.parametrize(
    "format,module,member",
    [
        ("docx", "docx", "word/document.xml"),
        ("xlsx", "openpyxl", "xl/workbook.xml"),
        ("pptx", "pptx", "ppt/presentation.xml"),
    ],
)
def test_malformed_xml_is_a_document_error(format, module, member):
    backend = pytest.importorskip(module)
    document = getattr(
        backend,
        {"docx": "Document", "xlsx": "Workbook", "pptx": "Presentation"}[format],
    )()
    original = io.BytesIO()
    document.save(original)
    corrupted = io.BytesIO()
    with (
        zipfile.ZipFile(original) as archive,
        zipfile.ZipFile(corrupted, "w") as output,
    ):
        for name in archive.namelist():
            output.writestr(name, b"<broken" if name == member else archive.read(name))
    with pytest.raises(DocumentParseError):
        inspect_document(
            SourceSnapshot(Path(f"corrupt.{format}"), corrupted.getvalue())
        )


@pytest.mark.parametrize("owner", ["paragraph", "layout", "master"])
def test_pptx_inherits_effective_title_colors(owner, tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Inherited title"
    if owner == "paragraph":
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(
            12, 34, 56
        )
    elif owner == "layout":
        slide.slide_layout.placeholders[0].text_frame.paragraphs[
            0
        ].font.color.rgb = RGBColor(12, 34, 56)
    else:
        default = slide.slide_layout.slide_master._element.xpath(
            "p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr"
        )[0]
        for child in list(default):
            default.remove(child)
        fill, rgb = OxmlElement("a:solidFill"), OxmlElement("a:srgbClr")
        rgb.set("val", "0C2238")
        fill.append(rgb)
        default.append(fill)
    path = tmp_path / "inheritance.pptx"
    deck.save(path)
    report = inspect_document(path)
    assert report.slides[0].title_colors == ((12, 34, 56),)


def test_pptx_gradients_alpha_and_unknown_styles_preserve_facts(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.oxml.xmlchemy import OxmlElement

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Known text"
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = RGBColor(255, 255, 255)
    fill.gradient_stops[1].color.rgb = RGBColor(240, 241, 242)
    color = slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color
    color.rgb = RGBColor(0, 0, 0)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "0")
    color._color._xClr.append(alpha)
    path = tmp_path / "styles.pptx"
    deck.save(path)
    page = inspect_document(path).slides[0]
    assert page.background_colors == ((255, 255, 255), (240, 241, 242))
    assert page.title_colors == page.title_background_colors
    color.theme_color = MSO_THEME_COLOR.ACCENT_1
    color._color._xClr.set("val", "unknown")
    deck.save(path)
    report = inspect_document(path)
    assert "unresolved theme" in report.slides[0].title_error
    assert report.slides[0].background_error is None
    assert "Known text" in report.to_markdown()


def test_pptx_nested_shapes_table_spans_and_explicit_fallback_size(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    deck = pptx.Presentation()
    # Remove master/default font size so the caller's fallback is observable.
    for parent in (deck._element, deck.slide_masters[0]._element):
        for node in parent.iter():
            node.attrib.pop("sz", None)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    shape.text = "Grouped title"
    group.left = Inches(2)
    table = slide.shapes.add_table(1, 2, 0, Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged table"
    path = tmp_path / "groups.pptx"
    deck.save(path)
    page = inspect_document(
        path, pptx_options=PptxOptions(fallback_font_size=23)
    ).slides[0]
    assert len(page.shapes) == page.text_shape_count == 2
    assert page.shapes[0].bounds[0] == Inches(2)
    assert page.shapes[0].paragraphs[0].runs[0].font_size == 23
    assert len(page.shapes[1].table.cells) == 1
    assert page.shapes[1].table.cells[0].column_span == 2
